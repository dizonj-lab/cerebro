"""Word and PowerPoint extraction.

Both formats carry real structure — headings, tables, slides, speaker notes —
so it is preserved rather than flattened to a wall of text.
"""
import io
import zipfile

from app.services.artifacts.extraction.base import (
    Block,
    ExtractionFailed,
    ExtractionResult,
    SourceLocation,
)


class DocxAdapter:
    modality = "office"
    signature_modalities = ("container",)
    extensions = (".docx",)

    def extract(self, content: bytes, *, filename: str) -> ExtractionResult:
        import docx
        from docx.opc.exceptions import PackageNotFoundError

        try:
            document = docx.Document(io.BytesIO(content))
        except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as exc:
            raise ExtractionFailed(f"Could not read the Word document: {exc}") from exc

        blocks: list[Block] = []
        # The most recent heading becomes the section for following paragraphs.
        section: str | None = None
        headings = 0

        for index, paragraph in enumerate(document.paragraphs, start=1):
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name or "").lower() if paragraph.style else ""
            is_heading = style.startswith("heading") or style == "title"
            if is_heading:
                section = text
                headings += 1
            blocks.append(
                Block(
                    text=text,
                    kind="heading" if is_heading else "paragraph",
                    location=SourceLocation(
                        section=section, paragraph=index, char_start=0, char_end=len(text)
                    ),
                )
            )

        # Tables are rendered row-wise so cell relationships survive.
        for table_index, table in enumerate(document.tables, start=1):
            rows = [
                " | ".join(cell.text.strip() for cell in row.cells)
                for row in table.rows
                if any(cell.text.strip() for cell in row.cells)
            ]
            if rows:
                text = "\n".join(rows)
                blocks.append(
                    Block(
                        text=text,
                        kind="table",
                        location=SourceLocation(
                            section=f"Table {table_index}", char_start=0, char_end=len(text)
                        ),
                    )
                )

        if not blocks:
            raise ExtractionFailed("The Word document contains no readable text.")

        return ExtractionResult(
            blocks=blocks,
            structure={
                "paragraphs": len(document.paragraphs),
                "headings": headings,
                "tables": len(document.tables),
            },
        )


class PptxAdapter:
    modality = "office"
    signature_modalities = ("container",)
    extensions = (".pptx",)

    def extract(self, content: bytes, *, filename: str) -> ExtractionResult:
        from pptx import Presentation
        from pptx.exc import PackageNotFoundError

        try:
            deck = Presentation(io.BytesIO(content))
        except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as exc:
            raise ExtractionFailed(f"Could not read the presentation: {exc}") from exc

        blocks: list[Block] = []
        notes_count = 0

        for slide_number, slide in enumerate(deck.slides, start=1):
            title = None
            if slide.shapes.title is not None:
                title = (slide.shapes.title.text or "").strip() or None

            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                text = shape.text_frame.text.strip()
                if not text:
                    continue
                blocks.append(
                    Block(
                        text=text,
                        kind="heading" if text == title else "body",
                        location=SourceLocation(
                            slide=slide_number, section=title,
                            char_start=0, char_end=len(text),
                        ),
                    )
                )

            # Speaker notes are often where the actual argument lives.
            if slide.has_notes_slide:
                notes = (slide.notes_slide.notes_text_frame.text or "").strip()
                if notes:
                    notes_count += 1
                    blocks.append(
                        Block(
                            text=notes,
                            kind="notes",
                            location=SourceLocation(
                                slide=slide_number, section=title,
                                char_start=0, char_end=len(notes),
                            ),
                        )
                    )

        if not blocks:
            raise ExtractionFailed("The presentation contains no readable text.")

        return ExtractionResult(
            blocks=blocks,
            structure={
                "slides": len(deck.slides),
                "slides_with_notes": notes_count,
                "blocks": len(blocks),
            },
        )
