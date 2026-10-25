"""Page furniture: lines a document repeats on most pages or slides are removed before fragments.

What matters: footers and page counters go, their first printing stays, real content that merely
looks alike ("Topic 1", "Topic 2", a column of numbers) is never touched, offsets still point at the
right characters, and the removal is reported. Uploads then neither propose a footer as a topic nor
carry it in every fragment.
"""

import io

import pytest
from sqlalchemy import select

from app.models import KnowledgeFragment
from app.services.artifacts import boilerplate
from app.services.artifacts.extraction.base import Block, ExtractionResult, SourceLocation
from tests.fixtures import make_pdf

FOOTER = "© 2026 National University of Singapore. All Rights Reserved"


def pages(texts: list[str], *, slide: bool = False) -> ExtractionResult:
    blocks = []
    for i, t in enumerate(texts, start=1):
        loc = SourceLocation(slide=i) if slide else SourceLocation(page=i)
        loc.char_end = len(t)
        blocks.append(Block(text=t, kind="paragraph", location=loc))
    return ExtractionResult(blocks=blocks)


def test_a_footer_on_every_page_keeps_only_its_first_printing():
    source = pages(
        [f"Topic {i}\nNeural networks part {i}.\n{FOOTER} Page {i}" for i in range(1, 6)]
    )
    cleaned, report = boilerplate.strip(source)

    text = cleaned.text
    assert text.count("All Rights Reserved") == 1
    for i in range(1, 6):
        assert f"Topic {i}" in text and f"Neural networks part {i}." in text
    assert (
        report["removed"] == 4
        and report["units"] == 5
        and report["method"] == "repeated_line_removal"
    )
    assert cleaned.structure["boilerplate"] == report
    # Every remaining block still points at exactly its own characters.
    originals = {b.location.page: b.text for b in source.blocks}
    for b in cleaned.blocks:
        assert originals[b.location.page][b.location.char_start : b.location.char_end] == b.text


def test_short_documents_and_documents_without_pages_are_untouched():
    two_pages = pages([f"Intro\n{FOOTER}", f"More\n{FOOTER}"])
    assert boilerplate.strip(two_pages) == (two_pages, None)

    word = ExtractionResult(
        blocks=[
            Block(text=f"Para {i}\n{FOOTER}", location=SourceLocation(paragraph=i))
            for i in range(5)
        ]
    )
    assert boilerplate.strip(word) == (word, None)


def test_numbers_and_mid_page_lines_on_a_pdf_are_not_furniture():
    body = "\n".join(f"Row {r} value" for r in range(12))
    source = pages([f"Heading {i}\n{body}\n42\n{FOOTER}" for i in range(1, 5)])
    cleaned, _ = boilerplate.strip(source)

    # "Row 5 value" repeats on every page, but in the middle of it: a table, not a footer.
    assert cleaned.text.count("Row 5 value") == 4
    # A bare number repeats too; with no letters it cannot be told from a table cell.
    assert cleaned.text.count("\n42") == 4
    assert cleaned.text.count(FOOTER) == 1


def test_a_line_repeated_anywhere_on_slides_is_furniture():
    source = pages(
        [f"Course: Pattern Recognition Systems\nSlide topic {i}\nDetail {i}" for i in range(1, 7)],
        slide=True,
    )
    cleaned, report = boilerplate.strip(source)
    assert cleaned.text.count("Course: Pattern Recognition Systems") == 1
    assert report["removed"] == 5


@pytest.mark.parametrize(
    "a, b, same",
    [
        (f"{FOOTER}Page 3", f"{FOOTER}Page 14", True),
        ("Slide 2 of 20", "Slide 9 of 20", True),
        ("Topic 1", "Topic 2", False),
        ("Lecture 3", "Lecture 4", False),
    ],
)
def test_only_page_counters_are_ignored_when_comparing_lines(a, b, same):
    assert (boilerplate.normalise(a) == boilerplate.normalise(b)) is same


# ---- through upload and registration


@pytest.fixture
def signed_in(client, valid_signup):
    client.post("/api/auth/signup", json=valid_signup)
    return valid_signup


def deck_with_footer(n: int = 5) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    for i in range(1, n + 1):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = f"Convolutional layers, part {i}"
        slide.placeholders[1].text = f"Filters and feature maps, example {i}."
        footer = slide.shapes.add_textbox(Inches(0.5), Inches(7), Inches(9), Inches(0.4))
        footer.text_frame.text = FOOTER
    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue()


def test_an_uploaded_deck_proposes_no_footer_topic_and_fragments_carry_it_once(
    client, signed_in, db
):
    pptx = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    a = client.post(
        "/api/artifacts", files={"file": ("deck.pptx", deck_with_footer(), pptx)}
    ).json()
    assert a["extraction"]["structure"]["boilerplate"]["removed"] == 4
    drafted = " ".join(a["metadata"].get("topics", []) + a["metadata"].get("tags", [])).lower()
    assert "rights reserved" not in drafted and "national university" not in drafted

    r = client.patch(f"/api/artifacts/{a['id']}", json={"metadata": a["metadata"], "approve": True})
    assert r.status_code == 200, r.text
    contents = db.scalars(
        select(KnowledgeFragment.content).where(KnowledgeFragment.artifact_id == a["id"])
    ).all()
    assert sum(c.count("All Rights Reserved") for c in contents) == 1
    assert sum(c.count("Convolutional layers, part") for c in contents) == 5


def test_an_uploaded_pdf_loses_its_footer_but_not_its_pages(client, signed_in):
    pdf = make_pdf(
        [
            f"Recurrent networks, session {i}\nGates control memory {i}.\n{FOOTER}"
            for i in range(1, 5)
        ]
    )
    a = client.post("/api/artifacts", files={"file": ("rnn.pdf", pdf, "application/pdf")}).json()
    report = a["extraction"]["structure"]["boilerplate"]
    assert report["removed"] == 3 and report["units"] == 4
    assert a["extraction"]["structure"]["pages"] == 4
