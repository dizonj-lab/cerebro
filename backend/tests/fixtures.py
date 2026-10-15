"""Small representative documents, generated rather than committed as blobs."""
import io


def make_pdf(pages: list[str]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    for text in pages:
        y = 800
        for line in text.split("\n"):
            c.drawString(60, y, line)
            y -= 16
        c.showPage()
    c.save()
    return buf.getvalue()


def make_docx(*, title: str, paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    import docx

    d = docx.Document()
    d.add_heading(title, level=1)
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for col, value in enumerate(row):
                t.cell(r, col).text = value
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_pptx(slides: list[tuple[str, str, str | None]]) -> bytes:
    from pptx import Presentation

    deck = Presentation()
    layout = deck.slide_layouts[1]
    for title, body, notes in slides:
        slide = deck.slides.add_slide(layout)
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue()
